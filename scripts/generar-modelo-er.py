"""Modelo lógico documental. No crea base de datos ni migraciones.

Ejecutar con Python del runtime Codex. Las entidades y relaciones son la fuente
única de DOCX, PDF, Markdown, Mermaid, SVG y draw.io. No usa servicios externos.
"""
from pathlib import Path
import json, re, math, html, textwrap, zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics.shapes import Drawing, Rect, Line, String
from reportlab.graphics import renderPDF, renderSVG
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.colors import HexColor
from PIL import Image, ImageDraw, ImageFont
from rutas_modelo_er import carpeta_lamina

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/arquitectura/modelo-er'
TMP=ROOT/'tmp/modelo-er'
OUT.mkdir(parents=True,exist_ok=True)
TMP.mkdir(parents=True,exist_ok=True)
(ROOT/'output/docx').mkdir(parents=True,exist_ok=True)
(ROOT/'output/pdf').mkdir(parents=True,exist_ok=True)
E={}
GROUPS={}

def entity(name, group, scope, purpose, source, fields, rules='', status='Propuesta lógica de requisito vigente'):
    fs=[]
    if scope=='T': fs.append(dict(name='tenant_id',type='uuid',null=False,key='FK',ref='ENTIDAD',note='UUID de la entidad autorizada; parte de toda FK local; sin generación propia'))
    fs.append(dict(name='id',type='uuid',null=False,key='PK',ref=None,note='UUIDv7 interno estable; generación en PostgreSQL con DEFAULT uuidv7(); ADR-015'))
    for item in fields.split(';'):
        parts=item.strip().split('|')
        if len(parts)<2: continue
        n,t=parts[:2]; note=parts[2] if len(parts)>2 else ''
        nullable=n.endswith('?'); n=n.rstrip('?')
        ref=t[1:] if t.startswith('@') else None
        fs.append(dict(name=n,type='uuid' if ref else t,null=nullable,key='FK' if ref else '',ref=ref,note=note))
    E[name]=dict(name=name,group=group,scope=scope,purpose=purpose,source=source,fields=fs,rules=rules,status=status)
    GROUPS.setdefault(group,[]).append(name)

# G global restringido; T entidad; X operación excepcional multientidad.
entity('CUENTA','01 Identidad y credenciales','G','Identidad personal reutilizable, separada de la ficha y los roles.','R09.a R10.a; ADR-012; docs/20',
 'nombre|texto;email_acceso?|texto;email_verificado_en?|instante;estado|enum;creada_en|instante',
 'No se fusiona por correo o documento. La baja institucional conserva la cuenta. Normalización y unicidad de contacto por revisar.')
entity('IDENTIDAD_EXTERNA','01 Identidad y credenciales','G','Identidad estable del proveedor vinculada mediante verificación.','R09.a; docs/12',
 'cuenta_id|@CUENTA;emisor|texto;sujeto|texto;vinculada_en|instante;revocada_en?|instante',
 'UQ(emisor,sujeto). El correo del proveedor no reemplaza esta clave.')
entity('CREDENCIAL_LOCAL','01 Identidad y credenciales','G','Credencial protegida de una cuenta que admite contraseña.','R10.a; S01; docs/20',
 'cuenta_id|@CUENTA;hash_contrasena|secreto;algoritmo_version|texto;actualizada_en|instante;revocada_en?|instante',
 'Una credencial local vigente por cuenta. No crear contraseña a cuentas solo externas mediante recuperación.')
entity('SESION','01 Identidad y credenciales','G','Sesión y revocación independiente de membresías y habilitaciones.','S01 S04; D16',
 'cuenta_id|@CUENTA;huella_sesion|secreto;iniciada_en|instante;vence_en|instante;revocada_en?|instante;mfa_verificada_en?|instante',
 'La sesión no almacena permisos permanentes ni sustituye su revalidación. Mecanismo por definir.')
entity('FACTOR_MFA','01 Identidad y credenciales','G','Segundo factor administrativo y ciclo de revocación.','S06; docs/20; D16',
 'cuenta_id|@CUENTA;tipo|enum;referencia_secreta|secreto;verificado_en?|instante;revocado_en?|instante',
 'Recuperación de contraseña no retira MFA. Tipo y custodia por definir.')
entity('TOKEN_IDENTIDAD','01 Identidad y credenciales','G','Acción limitada de recuperación, contacto o invitación de identidad.','R10.a; docs/20',
 'cuenta_id?|@CUENTA;destinatario|texto;huella|secreto;proposito|enum;vence_en|instante;usado_en?|instante;revocado_en?|instante',
 'UQ(huella). Uso único con consumo atómico; GET no consume. Sin token en claro. No sirve para confirmar turnos.')
entity('FACULTAD_PLATAFORMA','01 Identidad y credenciales','G','Concesión global acotada para incorporación y operación.','R10; MP-35 MP-38; D03',
 'cuenta_id|@CUENTA;permiso_id|@PERMISO;otorgante_id|@CUENTA;desde|instante;hasta?|instante;revocada_en?|instante',
 'No concede lectura general de fichas ni suplantación.')
entity('CANAL_IDENTIDAD','01 Identidad y credenciales','G','Remitente de plataforma para recuperar una cuenta sin contexto de entidad.','R10.a; docs/20; D09 D16',
 'proveedor|texto;remitente|texto;secreto_ref|secreto;estado|enum',
 'Canal restringido de identidad, separado de los remitentes de avisos de pacientes. No se elige una clínica para recuperar la cuenta.')
entity('TAREA_IDENTIDAD','01 Identidad y credenciales','G','Trabajo persistido de recuperación o verificación de identidad.','R10.a; docs/20',
 'token_id|@TOKEN_IDENTIDAD;canal_id|@CANAL_IDENTIDAD;tipo|enum;clave_logica|texto;programada_en|instante;estado|enum;intentos|entero',
 'UQ(clave_logica). Sin datos de pacientes ni membresías; contenido mínimo y sin tokens en logs. Respuesta pública no revela cuenta.')
entity('ENVIO_IDENTIDAD','01 Identidad y credenciales','G','Intento técnico del correo de identidad con entrega y fallo.','R10.a; docs/20',
 'tarea_id|@TAREA_IDENTIDAD;numero|entero;referencia_proveedor?|texto;iniciado_en|instante;estado|enum;error_codigo?|texto',
 'UQ(tarea_id,numero). Se revalida token y presupuesto global; no reutiliza canal ni permisos de una entidad ajena.')

entity('ENTIDAD','02 Organización y permisos','G','Cliente y raíz del aislamiento, incluida actividad independiente.','R01 R01.a; ADR-001',
 'tipo|enum;nombre|texto;zona_horaria|texto;estado|enum;creada_en|instante',
 'Cada fila es una raíz de tenant. No cambiar una clínica completa por la entidad de un profesional.')
entity('SEDE','02 Organización y permisos','T','Lugar de atención dentro de una entidad.','R01 R06',
 'nombre|texto;direccion?|texto;zona_horaria|texto;estado|enum')
entity('MEMBRESIA','02 Organización y permisos','T','Pertenencia del equipo con vigencia local.','R04 R10; docs/20',
 'cuenta_id|@CUENTA;desde|instante;hasta?|instante;estado|enum',
 'UQ(tenant_id,cuenta_id). Reactivar preserva la identidad local; historial en auditoría.')
entity('ROL','02 Organización y permisos','G','Definición de agrupaciones de facultades sin concederlas.','R04 R10; docs/22',
 'codigo|texto;nombre|texto;ambito|enum', 'UQ(codigo). Catálogo técnico propuesto; no asigna permisos pendientes.')
entity('PERMISO','02 Organización y permisos','G','Acción autorizable y tipo de ámbito.','R04 R12; docs/22; D03',
 'codigo|texto;accion|texto;tipo_ambito|enum', 'UQ(codigo). Códigos MP son trazabilidad y no permisos definitivos.')
entity('ROL_PERMISO','02 Organización y permisos','G','Descompone roles en acciones explícitas.','R04; docs/22',
 'rol_id|@ROL;permiso_id|@PERMISO', 'UQ(rol_id,permiso_id). Contenido a aprobar; sin concesiones automáticas por cargo.')
entity('MEMBRESIA_ROL','02 Organización y permisos','T','Asignación múltiple de roles al miembro de la entidad.','R04 R10',
 'membresia_id|@MEMBRESIA;rol_id|@ROL;otorgante_id|@CUENTA;desde|instante;hasta?|instante;revocada_en?|instante')
entity('MEMBRESIA_PERMISO','02 Organización y permisos','T','Facultades locales expresas adicionales y vigencia.','R04 R06; D03',
 'membresia_id|@MEMBRESIA;permiso_id|@PERMISO;otorgante_id|@CUENTA;desde|instante;hasta?|instante;revocada_en?|instante',
 'No sustituye delegación de agenda ni participación de chat.')

entity('PROFESIONAL_GLOBAL','03 Profesionales e incorporación','G','Identidad profesional verificada para coordinar capacidad.','R03 R05; ADR-011; D02',
 'cuenta_id?|@CUENTA;nombre_referencia|texto;matricula_referencia?|texto;verificacion_ref?|texto;estado|enum',
 'UQ(cuenta_id) cuando presente. No deduplicar por nombre. Referencia histórica mínima tras eliminación por resolver.')
entity('PROFESIONAL_ENTIDAD','03 Profesionales e incorporación','T','Vínculo profesional local, separado de la membresía de acceso.','R01 R05; ADR-011',
 'profesional_id|@PROFESIONAL_GLOBAL;membresia_id?|@MEMBRESIA;desde|instante;hasta?|instante;estado|enum',
 'UQ(tenant_id,profesional_id). Finalizar no libera turnos ni destruye fichaje. Retención de referencia histórica por revisar.')
entity('SOLICITUD_ENTIDAD','03 Profesionales e incorporación','G','Trámite de incorporación del solicitante y designación administrativa.','docs/20; CU-35; D14',
 'solicitante_id|@CUENTA;tipo_entidad|enum;nombre_propuesto|texto;relacion_declarada|texto;admin_contacto?|texto;estado|enum;entidad_resultado_id?|@ENTIDAD;revisor_id?|@CUENTA;resuelta_en?|instante',
 'No concede administración, contrato ni suscripción.', 'Propuesta funcional no aprobada')
entity('INVITACION_EQUIPO','03 Profesionales e incorporación','T','Invitación personal al equipo con alcance predeterminado.','docs/20; CU-07; D14',
 'emisor_id|@CUENTA;destinatario|texto;cuenta_destino_id?|@CUENTA;token_id|@TOKEN_IDENTIDAD;rol_propuesto_id?|@ROL;estado|enum;vence_en|instante;membresia_id?|@MEMBRESIA',
 'UQ(token_id). Correspondencia verificada al aceptar; invitación no delega agendas.')

entity('PACIENTE','04 Ficha y acceso del paciente','T','Ficha administrativa local, exista o no cuenta.','R01.a R09; ADR-010; docs/18',
 'nombres|texto;apellidos|texto;tipo_documento?|texto;documento_declarado?|texto;nacimiento?|fecha;domicilio?|texto;estado|enum;estado_datos_cobertura|enum',
 'Campos concretos propuestos. No guarda edad redundante ni historia clínica; sin dato no significa sin cobertura.')
entity('CONTACTO_PACIENTE','04 Ficha y acceso del paciente','T','Contacto administrativo que no otorga representación.','R09 R11; docs/18',
 'paciente_id|@PACIENTE;tipo|enum;valor|texto;titular_descripcion?|texto;verificado_en?|instante;estado|enum',
 'Teléfono familiar no acredita identidad ni es único por paciente.')
entity('SOLICITUD_ADMISION','04 Ficha y acceso del paciente','T','Solicitud protegida que nunca ocupa capacidad.','R09.a; ADR-007; docs/11-12',
 'solicitante_id?|@CUENTA;profesional_entidad_id?|@PROFESIONAL_ENTIDAD;datos_declarados|json;contacto_verificado_en?|instante;paciente_resuelto_id?|@PACIENTE;estado|enum;revisor_id?|@CUENTA;motivo_interno?|texto;resuelta_en?|instante',
 'Solicitante puede actuar por otro; sin cuenta previa es posible. No fusionar por coincidencias.')
entity('REPRESENTACION','04 Ficha y acceso del paciente','T','Relación verificada entre representante y paciente.','R09; D08',
 'paciente_id|@PACIENTE;representante_id|@CUENTA;tipo|enum;evidencia_ref|texto;verificador_id|@CUENTA;desde|instante;hasta?|instante;restricciones?|texto;revocada_en?|instante',
 'Varios representantes posibles. Contacto de emergencia y titular de cobertura no confieren esta relación.')
entity('REPRESENTACION_PERMISO','04 Ficha y acceso del paciente','T','Facultad concreta de representación.','R09; docs/22; D08',
 'representacion_id|@REPRESENTACION;permiso_id|@PERMISO', 'UQ(tenant_id,representacion_id,permiso_id). Gestión de turnos y lectura de chat son distintas.')
entity('ACCESO_PACIENTE','04 Ficha y acceso del paciente','T','Cuenta vinculada a ficha con admisión y activación verificadas.','R09.a; ADR-007/008',
 'paciente_id|@PACIENTE;cuenta_id|@CUENTA;representacion_id?|@REPRESENTACION;solicitud_id?|@SOLICITUD_ADMISION;aprobador_id|@CUENTA;aprobado_en|instante;activado_en?|instante;estado|enum;vinculacion_ref|texto',
 'UQ(tenant_id,paciente_id,cuenta_id). Sin representación opera como propio tras verificar correspondencia; con representación se exige paciente y cuenta coincidentes.')
entity('ALCANCE_ACCESO','04 Ficha y acceso del paciente','T','Contextos concretos autorizados del portal.','R09.a; D14',
 'acceso_id|@ACCESO_PACIENTE;profesional_entidad_id|@PROFESIONAL_ENTIDAD;agenda_id?|@AGENDA;permiso_id|@PERMISO;desde|instante;hasta?|instante;revocado_en?|instante',
 'Profesional obligatorio; agenda nula propone alcance a sus agendas locales. Semántica exacta pendiente; no habilita otras entidades.')
entity('PACIENTE_PROFESIONAL','04 Ficha y acceso del paciente','T','Relación numerada, independiente de agendas y cuenta.','R01.a; ADR-001',
 'profesional_id|@PROFESIONAL_GLOBAL;profesional_entidad_id|@PROFESIONAL_ENTIDAD;paciente_id|@PACIENTE;numero_visible|entero;estado|enum;alta_en|instante',
 'UQ(tenant_id,profesional_id,paciente_id) y UQ(tenant_id,profesional_id,numero_visible). Número positivo; profesional coincide con vínculo local; no renumerar ni reutilizar.')
entity('CONTADOR_PACIENTE','04 Ficha y acceso del paciente','T','Asignación transaccional por serie profesional y entidad.','R01.a; ADR-001',
 'profesional_id|@PROFESIONAL_GLOBAL;profesional_entidad_id|@PROFESIONAL_ENTIDAD;ultimo_confirmado|entero;version|entero',
 'UQ(tenant_id,profesional_id) y UQ(tenant_id,profesional_entidad_id). Inicial 0; alta serializada e idempotente; no usar COUNT o MAX sin serialización.')

entity('AGENDA','05 Agendas y catálogo','T','Agenda de un vínculo profesional y sede.','R04 R06',
 'profesional_entidad_id|@PROFESIONAL_ENTIDAD;sede_id|@SEDE;nombre|texto;zona_horaria|texto;estado|enum', 'Sede y vínculo profesional pertenecen al mismo tenant.')
entity('DELEGACION_AGENDA','05 Agendas y catálogo','T','Autorización por persona y agenda otorgada por su profesional.','R04',
 'agenda_id|@AGENDA;membresia_id|@MEMBRESIA;otorgante_profesional_id|@PROFESIONAL_ENTIDAD;desde|instante;hasta?|instante;revocada_en?|instante',
 'Otorgante debe ser titular de la agenda; destinatario es miembro activo de la misma entidad. Sin subdelegación implícita.')
entity('DELEGACION_PERMISO','05 Agendas y catálogo','T','Acciones específicas delegadas.','R04; D03',
 'delegacion_id|@DELEGACION_AGENDA;permiso_id|@PERMISO', 'UQ(tenant_id,delegacion_id,permiso_id). No concede chat ni admisión por operar agenda.')
entity('PRACTICA','05 Agendas y catálogo','T','Identidad estable de una práctica administrativa de la entidad.','R06; ADR-009; docs/16',
 'codigo_local|texto;origen|enum;base_inicial?|enum;estado|enum',
 'UQ(tenant_id,codigo_local). UQ parcial(tenant_id,base_inicial). Base CONSULTA reservada al alta inicial; personalizada exige base nula e inmutable.')
entity('VERSION_PRACTICA','05 Agendas y catálogo','T','Definición inmutable usada por las ofertas.','R06 R12; docs/16',
 'practica_id|@PRACTICA;numero_version|entero;nombre|texto;clasificacion|texto;descripcion?|texto;autor_id|@CUENTA;creada_en|instante',
 'UQ(tenant_id,practica_id,numero_version). Cambiar nombre o clasificación no cambia base ni permite autoselección.')
entity('OFERTA','05 Agendas y catálogo','T','Habilitación de una práctica en una agenda concreta.','R06; docs/16',
 'agenda_id|@AGENDA;practica_id|@PRACTICA;estado|enum', 'UQ(tenant_id,agenda_id,practica_id). Crear catálogo no habilita ofertas.')
entity('VERSION_OFERTA','05 Agendas y catálogo','T','Condiciones inmutables y vigencia de la oferta.','R06 R07; docs/16',
 'oferta_id|@OFERTA;version_practica_id|@VERSION_PRACTICA;numero_version|entero;duracion_min|entero;preparacion_min|entero;cierre_min|entero;modalidad|enum;portal_habilitado|bool;gestion_personal|bool;particular_habilitado|bool;desde|instante;hasta?|instante',
 'UQ(tenant_id,oferta_id,numero_version). Duración >0 y pausas >=0. Práctica de versión coincide con oferta; portal implica base CONSULTA y no gestión obligatoria. Vigencias activas sin solapamiento.')
entity('HORARIO_AGENDA','05 Agendas y catálogo','T','Disponibilidad semanal recurrente según zona de agenda.','R06; D05',
 'agenda_id|@AGENDA;dia_semana|entero;hora_inicio|hora;hora_fin|hora;desde|fecha;hasta?|fecha',
 'Inicio anterior a fin; horarios nocturnos se dividen. No es ocupación ni garantiza disponibilidad efectiva.')

entity('EXCEPCION_AGENDA','06 Excepciones y recursos','T','Apertura o indisponibilidad puntual que no cancela turnos.','R03 R06; ADR-013',
 'agenda_id|@AGENDA;inicio|instante;fin|instante;tipo|enum;motivo_interno?|texto;lote_id?|@LOTE_REPROGRAMACION;autor_id|@CUENTA',
 'Inicio < fin. Bloquear una ausencia no libera los turnos ya existentes.')
entity('RECURSO','06 Excepciones y recursos','T','Recurso físico administrativo con disponibilidad a validar.','R03 R06; D02 D05',
 'sede_id|@SEDE;nombre|texto;tipo|enum;capacidad|entero;estado|enum', 'Control automático de recursos condicionado al piloto; no asumir sobrecupos.', 'Extensión propuesta pendiente de validación')
entity('OFERTA_RECURSO','06 Excepciones y recursos','T','Recursos exigidos por una versión de oferta.','R06; docs/16; D02',
 'version_oferta_id|@VERSION_OFERTA;recurso_id|@RECURSO;cantidad|entero;incluye_preparacion|bool;incluye_cierre|bool',
 'UQ(tenant_id,version_oferta_id,recurso_id). Sede del recurso coincide con agenda, salvo regla futura explícita.', 'Extensión propuesta pendiente de validación')
entity('RESERVA_RECURSO','06 Excepciones y recursos','T','Ocupación física vinculada al mismo bloqueo profesional.','R03 R06; D02',
 'recurso_id|@RECURSO;ocupacion_id|@OCUPACION_GLOBAL;inicio|instante;fin|instante;cantidad|entero;estado|enum;vence_en?|instante',
 'Capacidad del recurso y profesional se adquieren/liberan en la misma unidad consistente; alcance por aprobar.', 'Extensión propuesta pendiente de validación')

entity('COBERTURA_MEDICA','07 Coberturas y convenios','T','Catálogo local de obras sociales y prepagas, separado del cliente SaaS.','R01 R07; docs/18',
 'nombre|texto;codigo_local?|texto;estado|enum', 'Se propone catálogo local para evitar compartir configuración; un catálogo de referencia global futuro no compartiría afiliaciones.')
entity('PLAN_COBERTURA','07 Coberturas y convenios','T','Plan o producto de una cobertura médica local.','R07; D06',
 'cobertura_id|@COBERTURA_MEDICA;nombre|texto;codigo?|texto;estado|enum')
entity('AFILIACION_PACIENTE','07 Coberturas y convenios','T','Cobertura declarada y verificación administrativa diferenciadas.','R07; docs/18',
 'paciente_id|@PACIENTE;cobertura_id|@COBERTURA_MEDICA;plan_id?|@PLAN_COBERTURA;numero_afiliado?|texto;titular_declarado?|texto;parentesco?|texto;desde?|fecha;hasta?|fecha;estado_verificacion|enum;verificacion_ref?|texto',
 'Plan debe pertenecer a cobertura. No hacer único el número familiar de afiliado ni inferir representación. Pueden existir varias afiliaciones por paciente.')
entity('CONVENIO','07 Coberturas y convenios','T','Acuerdo aplicable a una oferta concreta, cobertura y plan.','R07; D06',
 'oferta_id|@OFERTA;cobertura_id|@COBERTURA_MEDICA;plan_id?|@PLAN_COBERTURA;estado|enum',
 'Granularidad explícita por oferta identifica práctica, agenda y profesional; plan nulo propone todos los planes de esa cobertura, sujeto a revisión y precedencia explícita.')
entity('VERSION_CONVENIO','07 Coberturas y convenios','T','Regla de cobertura inmutable que conserva vigencia y condiciones.','R07 R12',
 'convenio_id|@CONVENIO;numero_version|entero;desde|instante;hasta?|instante;requiere_autorizacion|bool;resultado_regla|enum;condiciones|json',
 'UQ(tenant_id,convenio_id,numero_version). Conflictos de reglas se resuelven explícitamente; no elegir particular automáticamente.')
entity('AUTORIZACION_COBERTURA','07 Coberturas y convenios','T','Constancia administrativa específica para una práctica.','R07; docs/18; D06',
 'afiliacion_id|@AFILIACION_PACIENTE;practica_id|@PRACTICA;referencia|texto;estado|enum;desde?|instante;hasta?|instante;verificador_id?|@CUENTA',
 'No equivale a afiliación verificada ni permiso de sistema. Sin API externa aprobada.')
entity('EVALUACION_ELEGIBILIDAD','07 Coberturas y convenios','T','Resultado y regla aplicados antes de adquirir y completar una reserva.','R02 R07 R12',
 'paciente_id|@PACIENTE;version_oferta_id|@VERSION_OFERTA;modalidad_elegida|enum;afiliacion_id?|@AFILIACION_PACIENTE;version_convenio_id?|@VERSION_CONVENIO;autorizacion_id?|@AUTORIZACION_COBERTURA;resultado|enum;evaluada_en|instante;evidencia|json',
 'Elegible, incompatible o revisión. Particular exige elección expresa y oferta habilitada. Cobertura exige afiliación del paciente y compatibilidad de plan/práctica/oferta.')

entity('OCUPACION_GLOBAL','08 Reservas y capacidad','G','Registro mínimo compartido para excluir intervalos del profesional.','R03; ADR-001',
 'profesional_id|@PROFESIONAL_GLOBAL;inicio|instante;fin|instante;estado|enum;vence_en?|instante;referencia_opaca|texto',
 'UQ(referencia_opaca). Sin paciente, tenant, agenda ni motivo. Inicio < fin; rangos [inicio,fin). Exclusión por profesional para ocupaciones activas. Vencimiento se resuelve explícitamente y transaccionalmente; sin índice dinámico now().')
entity('PRE_RESERVA','08 Reservas y capacidad','T','Bloqueo persistido previo al registro del turno.','R02 R03; CU-36',
 'paciente_profesional_id|@PACIENTE_PROFESIONAL;agenda_id|@AGENDA;version_oferta_id|@VERSION_OFERTA;elegibilidad_id|@EVALUACION_ELEGIBILIDAD;ocupacion_id|@OCUPACION_GLOBAL;actor_id|@CUENTA;representacion_id?|@REPRESENTACION;inicio_atencion|instante;fin_atencion|instante;vence_en|instante;estado|enum;renovaciones|entero',
 'UQ(ocupacion_id) dentro de esta tabla. Profesional, paciente y oferta coinciden. TTL pendiente; expirar no desplaza ocupaciones posteriores. No conservar transacción abierta durante el formulario.')
entity('TURNO','08 Reservas y capacidad','T','Identidad estable de una cita registrada para un paciente.','R02 R08 R12; CU-37',
 'paciente_profesional_id|@PACIENTE_PROFESIONAL;pre_reserva_id|@PRE_RESERVA;creador_id|@CUENTA;creado_en|instante;estado_turno|enum;estado_asistencia|enum;estado_jornada|enum;version_concurrencia|entero',
 'UQ(tenant_id,pre_reserva_id). Registro no confirma asistencia. Dimensiones separadas y combinaciones válidas por revisar. Una pre-reserva produce como máximo un turno.')
entity('VERSION_TURNO','08 Reservas y capacidad','T','Condiciones y horario históricos inmutables por registro o cambio.','R06 R07 R08 R12',
 'turno_id|@TURNO;numero_version|entero;agenda_id|@AGENDA;version_oferta_id|@VERSION_OFERTA;elegibilidad_id|@EVALUACION_ELEGIBILIDAD;ocupacion_id|@OCUPACION_GLOBAL;inicio_atencion|instante;fin_atencion|instante;zona_horaria|texto;condiciones_aplicadas|json;creada_en|instante;sustituida_en?|instante',
 'UQ(tenant_id,turno_id,numero_version). Una versión vigente por turno; snapshot conserva nombre, duración, pausas y cobertura. Primera versión reutiliza ocupación de pre-reserva; cambiar horario adquiere destino y retira origen atómicamente.')
entity('CONFIRMACION_ASISTENCIA','08 Reservas y capacidad','T','Acción explícita sobre una versión determinada.','R08 R09',
 'version_turno_id|@VERSION_TURNO;actor_id|@CUENTA;representacion_id?|@REPRESENTACION;medio|enum;confirmada_en|instante;referencia_evidencia?|texto',
 'UQ(tenant_id,version_turno_id). Repetir es idempotente; confirmar no reserva otra capacidad. Medio telefónico conserva operador y evidencia, con facultad específica.')
entity('EVENTO_TURNO','08 Reservas y capacidad','T','Historia cronológica de cambios, llegada, atención y cancelación.','R08 R12',
 'turno_id|@TURNO;version_turno_id?|@VERSION_TURNO;actor_id?|@CUENTA;servicio_actor?|texto;tipo|enum;ocurrido_en|instante;motivo?|texto;resultado|enum;cambios_minimos|json',
 'Actor humano o servicio identificados. La baja del profesional no borra eventos ni crea cancelación implícita.')
entity('IDEMPOTENCIA','08 Reservas y capacidad','T','Resultado de una operación repetible y su huella de entrada.','R02 R03 R12',
 'actor_id|@CUENTA;operacion|texto;clave|texto;huella_entrada|texto;resultado_ref|texto;estado|enum;creada_en|instante;vence_en?|instante',
 'UQ(tenant_id,actor_id,operacion,clave). Misma clave con entrada distinta se rechaza. No almacena secretos; retención por definir.')

entity('LOTE_REPROGRAMACION','09 Reprogramación y gestión','T','Operación del profesional por ausencia en agendas autorizadas.','R11.b; ADR-013',
 'profesional_entidad_id|@PROFESIONAL_ENTIDAD;iniciador_id|@CUENTA;motivo_interno|texto;estado|enum;creado_en|instante;clave_operacion|texto',
 'UQ(tenant_id,clave_operacion). WhatsApp a todos los afectados, incluso confirmados. Estados y política de liberación pendientes.')
entity('INTERVALO_LOTE','09 Reprogramación y gestión','T','Uno o más tramos afectados y agenda local.','R11.b; ADR-013',
 'lote_id|@LOTE_REPROGRAMACION;agenda_id|@AGENDA;inicio|instante;fin|instante',
 'Inicio < fin y agenda del profesional del lote. Una fecha completa se convierte desde zona horaria de agenda. Detectar toda ocupación superpuesta, aunque el turno comience antes.')
entity('TURNO_AFECTADO','09 Reprogramación y gestión','T','Seguimiento separado del aviso y del reemplazo.','R11.b; ADR-013; docs/04',
 'lote_id|@LOTE_REPROGRAMACION;version_original_id|@VERSION_TURNO;turno_reemplazo_id?|@TURNO;estado_aviso|enum;estado_resolucion|enum;responsable_id?|@CUENTA;resuelto_en?|instante;motivo_cierre?|texto',
 'UQ(tenant_id,lote_id,version_original_id). Reemplazo no duplicado ni cruzado de paciente. Envío no resuelve reemplazo. Diseño propone nuevo turno para reemplazo de lote; cambio individual crea versión del mismo turno.')
entity('SOLICITUD_GESTION','09 Reprogramación y gestión','T','Pedido humano sin selección ni bloqueo de horarios.','R06 R11.a; CU-16 CU-26',
 'paciente_id|@PACIENTE;profesional_entidad_id|@PROFESIONAL_ENTIDAD;practica_id?|@PRACTICA;turno_afectado_id?|@TURNO_AFECTADO;solicitante_id|@CUENTA;responsable_id?|@CUENTA;motivo|texto;estado|enum;creada_en|instante;resuelta_en?|instante',
 'Un mensaje puede motivar este trámite; crearlo requiere acción explícita y nunca ocupa capacidad.')
entity('TOKEN_TURNO','09 Reprogramación y gestión','T','Token de acción limitado a una versión y destinatario autorizado.','R08; S07',
 'version_turno_id|@VERSION_TURNO;destino_id|@DESTINO_AVISO;huella|secreto;proposito|enum;vence_en|instante;usado_en?|instante;revocado_en?|instante',
 'UQ(huella). Cancelar/reprogramar invalida token obsoleto. GET no modifica; la acción exige autorización actual.')

entity('CANAL_ENTIDAD','10 Canales y destinatarios','T','Remitente y cuenta de proveedor verificados por entidad.','R11; docs/04',
 'canal|enum;proveedor|texto;cuenta_externa|texto;remitente|texto;secreto_ref|secreto;verificada_en?|instante;estado|enum',
 'Credencial solo referencia al almacén seguro. Correlación de webhooks por cuenta verificada, nunca por tenant recibido sin comprobar.')
entity('DESTINO_AVISO','10 Canales y destinatarios','T','Contacto autorizado para un paciente y propósito.','R09 R11; S09',
 'paciente_id|@PACIENTE;contacto_id|@CONTACTO_PACIENTE;cuenta_destinataria_id?|@CUENTA;representacion_id?|@REPRESENTACION;canal|enum;proposito|enum;autorizacion_ref|texto;desde|instante;hasta?|instante;revocado_en?|instante',
 'Contacto pertenece al paciente. Representación, si presente, coincide con paciente y cuenta. Autorización y contacto se revalidan al enviar; recepción no concede portal.')
entity('PRESUPUESTO_CANAL','10 Canales y destinatarios','T','Límite técnico de consumo y continuidad operativa.','R11; docs/04; D09 D15',
 'canal_id|@CANAL_ENTIDAD;periodo_inicio|instante;periodo_fin|instante;moneda|texto;limite_tecnico|decimal;umbral_alerta|decimal;politica_continuidad|texto',
 'No es cuota facturada ni suscripción; agotamiento no cancela turnos ni cierra portal.')
entity('CONSUMO_CANAL','10 Canales y destinatarios','T','Costo técnico estimado o conciliado por intento.','R11; docs/04',
 'intento_id|@INTENTO_ENVIO;moneda|texto;importe|decimal;tipo_medicion|enum;tarifa_version_ref|texto;registrado_en|instante',
 'Estimación y conciliación se distinguen; no sumar ambas como dos cargos del mismo intento.')

entity('TAREA_PERSISTIDA','11 Avisos y entrega','T','Trabajo diferido recuperable, consistente con la operación de origen.','R08 R11 R12; docs/04',
 'version_turno_id?|@VERSION_TURNO;turno_afectado_id?|@TURNO_AFECTADO;tipo|enum;deduplicacion|texto;programada_en|instante;estado|enum;intentos|entero;bloqueada_hasta?|instante;version_politica_ref|texto',
 'UQ(tenant_id,deduplicacion). Correo único dentro de T-72/T-48; WA T-36 solo pendiente; T-24 cierra. Reprogramación profesional es evento independiente. Trabajos sin turno llevan tipo y referencia de origen controlada; no datos arbitrarios.')
entity('AVISO_EXTERNO','11 Avisos y entrega','T','Mensaje lógico individual, separado de cada intento del proveedor.','R11 R11.b; ADR-013',
 'tarea_id|@TAREA_PERSISTIDA;destino_id|@DESTINO_AVISO;canal_id|@CANAL_ENTIDAD;clave_logica|texto;plantilla_version_ref|texto;estado|enum;creado_en|instante',
 'UQ(tenant_id,clave_logica). Dedupe por evento, versión, destinatario y canal; reprogramación exige WhatsApp. Fallo queda pendiente de seguimiento.')
entity('INTENTO_ENVIO','11 Avisos y entrega','T','Ejecución técnica del mensaje y correlación externa.','R11; docs/04',
 'aviso_id|@AVISO_EXTERNO;numero|entero;referencia_proveedor?|texto;iniciado_en|instante;estado|enum;error_codigo?|texto;entregado_en?|instante;leido_en?|instante',
 'UQ(tenant_id,aviso_id,numero). Ref externa única dentro de cuenta proveedor cuando exista. No prometer entrega exactamente una vez.')
entity('EVENTO_PROVEEDOR','11 Avisos y entrega','T','Webhook autenticado, correlacionado y deduplicado.','R11; S08',
 'canal_id|@CANAL_ENTIDAD;intento_id|@INTENTO_ENVIO;evento_externo|texto;tipo|enum;ocurrido_en|instante;recibido_en|instante;validacion_ref|texto',
 'UQ(tenant_id,canal_id,evento_externo). Canal coincide con aviso del intento. Lectura, entrega o texto libre no confirman asistencia.')
entity('AVISO_INTERNO','11 Avisos y entrega','T','Aviso persistido a una cuenta autorizada, incluido cierre T-24.','R08 R11.a R12; MP-33',
 'tarea_id?|@TAREA_PERSISTIDA;version_turno_id?|@VERSION_TURNO;conversacion_id?|@CONVERSACION;destinatario_id|@CUENTA;tipo|enum;clave_logica|texto;creado_en|instante;leido_en?|instante;resuelto_en?|instante',
 'UQ(tenant_id,clave_logica,destinatario_id). Consultar exige permiso actual aun tras recibir el aviso. Aviso y cierre T-24 se persisten consistentemente.')

entity('CONVERSACION','12 Chat humano','T','Conversación de paciente y profesional dentro de una entidad.','R11.a; D13',
 'paciente_profesional_id|@PACIENTE_PROFESIONAL;estado|enum;creada_en|instante;cerrada_en?|instante',
 'Número de conversaciones simultáneas por vínculo pendiente. No participa secretaría por delegación de agenda.')
entity('PARTICIPANTE_CHAT','12 Chat humano','T','Participación vigente y verificable de una cuenta.','R09 R11.a; S14',
 'conversacion_id|@CONVERSACION;cuenta_id|@CUENTA;representacion_id?|@REPRESENTACION;tipo|enum;desde|instante;hasta?|instante;revocado_en?|instante',
 'UQ(tenant_id,conversacion_id,cuenta_id). Cuenta debe corresponder al paciente/profesional o representante habilitado; participación administrativa pendiente.')
entity('PARTICIPANTE_PERMISO','12 Chat humano','T','Facultades de lectura, envío o respuesta independientes de agenda.','R11.a; docs/22; D13',
 'participante_id|@PARTICIPANTE_CHAT;permiso_id|@PERMISO', 'UQ(tenant_id,participante_id,permiso_id). Participar no concede todas las operaciones.')
entity('MENSAJE','12 Chat humano','T','Texto humano persistido con autoría real.','R11.a R12',
 'conversacion_id|@CONVERSACION;participante_autor_id|@PARTICIPANTE_CHAT;texto|texto;enviado_en|instante;estado|enum;correccion_de_id?|@MENSAJE',
 'Autor y corrección pertenecen a la conversación. Correcciones trazables por revisar; no adjuntos, IA ni cambios de turno por inferencia.')
entity('LECTURA_MENSAJE','12 Chat humano','T','Lectura por participante, separada de atención o respuesta.','R11.a; D13',
 'mensaje_id|@MENSAJE;participante_id|@PARTICIPANTE_CHAT;leido_en|instante',
 'UQ(tenant_id,mensaje_id,participante_id). Mensaje y participante de la misma conversación.')

entity('TRANSICION_PROFESIONAL','13 Continuidad del fichaje','X','Operación administrativa excepcional entre origen y destino.','R05 R10; ADR-011; docs/19',
 'profesional_id|@PROFESIONAL_GLOBAL;entidad_origen_id|@ENTIDAD;entidad_destino_id?|@ENTIDAD;operador_id|@CUENTA;tipo|enum;estado|enum;fecha_efectiva?|instante;autorizacion_ref|texto;clave_operacion|texto',
 'UQ(clave_operacion). Metadatos restringidos de plataforma; no almacena fichas. Baja puede no tener destino; continuidad sí lo exige. No modifica tenant de registros existentes.', 'Alternativa de modelado pendiente D04 D12')
entity('CORRESPONDENCIA_FICHAJE','13 Continuidad del fichaje','X','Procedencia explícita de vínculos origen y destino durante continuidad.','ADR-011; docs/19; D04 D12',
 'transicion_id|@TRANSICION_PROFESIONAL;vinculo_origen_id|@PACIENTE_PROFESIONAL;vinculo_destino_id?|@PACIENTE_PROFESIONAL;estado|enum;alcance_datos_ref|texto;resolucion_numeracion_ref?|texto',
 'Excepción cruzada explícita: validar tenant origen/destino contra transición y mismo profesional. UQ(transicion_id,vinculo_origen_id). No habilita lectura general ni decide copia/proyección. Números destino pendientes.', 'Alternativa de modelado pendiente D04 D12')
entity('PETICION_ELIMINACION','13 Continuidad del fichaje','G','Petición expresa del profesional diferenciada de ejecución.','R05; ADR-011; CU-31',
 'profesional_id|@PROFESIONAL_GLOBAL;solicitante_id?|@CUENTA;peticion_verificada_ref|texto;solicitada_en|instante;estado|enum;ejecutor_id?|@CUENTA;ejecutada_en?|instante;referencias_historicas_ref?|texto',
 'Solo petición verificada del propio profesional. Conservar fichas e historial; anonimización, retiro de vínculos y referencia histórica requieren diseño. Ningún borrado en cascada aprobado.')

entity('AUDITORIA_ENTIDAD','14 Auditoría y operación','T','Registro cronológico restringido del actor real y contexto local.','R12; S10',
 'actor_id?|@CUENTA;servicio_actor?|texto;paciente_id?|@PACIENTE;agenda_id?|@AGENDA;accion|texto;recurso_tipo|texto;recurso_ref|texto;ocurrido_en|instante;resultado|enum;motivo?|texto;cambios_minimos|json;correlacion|texto',
 'Actor humano o servicio identificados. Referencia polimórfica es localizador auditado, no FK; no secretos ni cuerpos de chat. No se borra por baja profesional.')
entity('AUDITORIA_GLOBAL','14 Auditoría y operación','G','Identidad y operación global sin exposición a administradores locales.','R10 R12; S10; docs/20',
 'actor_id?|@CUENTA;servicio_actor?|texto;cuenta_afectada_id?|@CUENTA;accion|texto;ocurrido_en|instante;resultado|enum;contexto_minimo|json',
 'Sin datos de ficha ni mensajes. Eventos de login, recuperación y facultades globales; acceso separado de auditoría local.')
entity('POLITICA_OPERATIVA','14 Auditoría y operación','T','Versión de parámetros de reserva, confirmación, cuotas y gasto.','R02.a R08 R11; D01 D07 D15',
 'tipo|enum;numero_version|entero;parametros|json;desde|instante;hasta?|instante;aprobador_id|@CUENTA',
 'UQ(tenant_id,tipo,numero_version). Sin valores inventados de TTL/cuotas. Tipo identifica esquema de parámetros; regla general T-24 no se anula silenciosamente.')
entity('EXCEPCION_POLITICA','14 Auditoría y operación','T','Excepción asistida, temporal y motivada a una política.','R02.a R08; D07 D15',
 'politica_id|@POLITICA_OPERATIVA;paciente_id?|@PACIENTE;turno_id?|@TURNO;actor_beneficiario_id?|@CUENTA;aprobador_id|@CUENTA;motivo|texto;desde|instante;hasta|instante;alcance|json',
 'Debe indicar al menos un sujeto/turno y alcance concreto; no es suspensión masiva ni permiso global.')
entity('CASO_ABUSO','14 Auditoría y operación','T','Revisión humana de señales sin compartir listas entre entidades.','R02.a; S16; CU-27',
 'paciente_id?|@PACIENTE;cuenta_id?|@CUENTA;revisor_id?|@CUENTA;senal_minima|json;estado|enum;decision?|texto;abierto_en|instante;cerrado_en?|instante',
 'DNI declarado o IP familiar no bastan para sanción. Métricas de frecuencia efímeras no requieren entidad de dominio por cada petición.')
entity('ACCESO_SOPORTE','14 Auditoría y operación','T','Concesión excepcional temporal a recursos delimitados.','R10; S02; MP-44; D03 D16',
 'operador_id|@CUENTA;concedente_id|@CUENTA;motivo|texto;alcance|json;desde|instante;hasta|instante;revocado_en?|instante',
 'Concedente, recursos y acciones aún por aprobar. No suplantación libre ni lectura por ser administrador de plataforma.', 'Propuesta de soporte pendiente')
entity('RESPALDO','14 Auditoría y operación','G','Evidencia operativa protegida del backup, sin credenciales ni fichas.','S11; CU-33',
 'operador_ref|texto;repositorio_seguro_ref|texto;iniciado_en|instante;finalizado_en?|instante;alcance_ref|texto;integridad_ref?|texto;estado|enum',
 'Metadatos de infraestructura fuera de API de negocio. Archivo real reside en almacenamiento protegido; incluye consistencia de turnos, contadores, ocupación y tareas.')
entity('ENSAYO_RESTAURACION','14 Auditoría y operación','G','Evidencia de recuperación en entorno separado.','S11; CU-33; D16',
 'respaldo_id|@RESPALDO;entorno_ref|texto;ejecutado_en|instante;duracion_seg|entero;perdida_datos_seg?|entero;resultado|enum;evidencia_ref|texto',
 'Sin envíos reales desde ensayo. RPO/RTO son objetivos por acordar, no valores garantizados.')
entity('INCIDENTE','14 Auditoría y operación','G','Coordinación técnica de contención y recuperación.','S12; CU-34',
 'responsable_ref|texto;abierto_en|instante;cerrado_en?|instante;estado|enum;alcance_minimo|texto;acciones_ref|texto;evidencia_ref?|texto',
 'Metadatos técnicos mínimos; información local sensible permanece bajo su entidad y acceso autorizado.')

TYPE_LABEL={'uuid':'UUID','texto':'Texto','enum':'Estado o código','instante':'Instante con zona','fecha':'Fecha','hora':'Hora local','entero':'Entero','bool':'Sí o no','decimal':'Decimal','json':'Estructura versionada','secreto':'Huella o referencia protegida'}

def relations():
    result=[]
    for e in E.values():
        for f in e['fields']:
            if not f['ref']: continue
            target=E[f['ref']]
            unique_single={('PROFESIONAL_GLOBAL','cuenta_id'),('INVITACION_EQUIPO','token_id'),('PRE_RESERVA','ocupacion_id'),('TURNO','pre_reserva_id'),('CONFIRMACION_ASISTENCIA','version_turno_id')}
            result.append(dict(child=e['name'],parent=f['ref'],field=f['name'],parent_card='0..1' if f['null'] else '1',child_card='0..1' if (e['name'],f['name']) in unique_single else '0..N',
              integrity='compuesta tenant_id + id' if e['scope']=='T' and target['scope']=='T' else 'excepción validada' if e['scope']=='X' else 'ID global restringido'))
    return result

def validate():
    for n,e in E.items():
        assert len({f['name'] for f in e['fields']})==len(e['fields']),n
        for f in e['fields']: assert not f['ref'] or f['ref'] in E,(n,f)
        assert (e['scope']=='T')==any(f['name']=='tenant_id' for f in e['fields']),n
    assert len(E)>70
validate()
R=relations()

NOTES=[
 ('Estado del modelo','Modelo lógico del alcance documentado al 16 de octubre de 2026, actualizado el 17 de octubre con ADR-015. Las reglas confirmadas se conservan. Se acepta UUID versión 7 para los IDs internos; las demás tablas, columnas, tipos, estados y mecanismos nuevos mantienen su condición de propuesta para revisión. No es una migración SQL ni acredita implementación.'),
 ('Identificadores internos','ADR-015 establece UUIDv7 para toda PK interna y tipo PostgreSQL uuid para PK y FK, incluido tenant_id. PostgreSQL genera cada PK mediante DEFAULT uuidv7(); las FK conservan el UUID del padre y no generan valores propios. El ID es inmutable, no se reutiliza y mantiene restricciones PK/UQ. Java usa java.util.UUID y JSON/frontend cadenas UUID con guiones. Los números visibles, contadores, versiones y códigos externos conservan su tipo propio.'),
 ('Uso de identificadores','El número visible del paciente es entero y pertenece al profesional dentro de la entidad; no se deriva del UUID ni lo sustituye. Las pantallas y reportes habituales usan ese número con su contexto. Conocer un UUID no concede acceso ni sustituye tokens secretos. UUIDv7 incorpora tiempo; no reemplaza creado_en, auditoría ni el orden de confirmación de transacciones. Se conservan tenant_id, claves compuestas, RLS y permisos del servidor.'),
 ('Ámbito y lectura','G significa global restringido; T pertenece a una entidad y lleva tenant_id obligatorio; X es una excepción administrativa entre origen y destino, accesible solo mediante operación controlada. Global no significa público. Cada ficha es local: no se crea un paciente global compartido por coincidencia de datos.'),
 ('Claves y cardinalidad','PK identifica la clave interna; FK identifica la referencia; UQ indica unicidad; ? significa opcional. Toda tabla T tiene PK(id) y UQ(tenant_id,id); cada referencia T a T usa FK(tenant_id,campo_id) a (tenant_id,id). Los IDs internos son únicos. Un padre admite 0..N hijos salvo UQ específico. La nulabilidad de la FK determina 1 o 0..1 padre por hijo; las UQ descritas pueden reducir los hijos a 0..1.'),
 ('Convención de atributos','Se muestran todos los atributos propuestos en el diccionario, JSON, Mermaid y draw.io. Para legibilidad, cada lámina enlaza las relaciones dentro de su grupo; las FK a otros grupos nombran su destino en la fila. El registro de relaciones enumera todas las referencias, incluido tenant_id. Restricciones de igualdad y ámbito adicionales se explicitan por entidad.'),
 ('Aislamiento y acceso','RLS y permisos del servidor complementan las claves compuestas. La cuenta, membresía, admisión, representación, alcance, delegación y participación de chat son comprobaciones separadas. Los roles no resuelven celdas pendientes de la matriz MP. La auditoría no concede acceso a los recursos referidos.'),
 ('Numeración estable','PACIENTE_PROFESIONAL conserva la serie por tenant_id y profesional_id. CONTADOR_PACIENTE empieza en 0 y asigna el primer número 1 dentro de una transacción serializada. La baja del 12 conserva su número, no altera el 13 y no permite reutilizar el 12. Reactivar devuelve la misma relación; no se infiere un total desde el último número.'),
 ('Reserva y capacidad','OCUPACION_GLOBAL no contiene pacientes, agendas ni entidades. PRE_RESERVA y primera VERSION_TURNO pueden referir a la misma ocupación, con el mismo profesional, pero solo una fase puede estar activa. La conversión conserva esa ocupación sin intervalo libre. Ninguna ocupación activa queda sin propietario local válido. Confirmar asistencia no crea ni libera capacidad.'),
 ('Horarios y concurrencia','Los intervalos son [inicio,fin), con inicio menor que fin. Se guardan instantes y zona de agenda; las recurrencias conservan hora local. La exclusión por profesional se aplica a ocupaciones activas; vencer requiere transición explícita. La política de locks, pausas de traslado y recursos sigue D02. Reprogramar adquiere destino y sustituye origen de forma atómica, conservando original si falla.'),
 ('Versiones y estados','VERSION_PRACTICA, VERSION_OFERTA, VERSION_CONVENIO y VERSION_TURNO son inmutables. Una sola versión de turno está vigente. Estados de admisión, pre-reserva, turno, asistencia, jornada, entrega y reemplazo son dimensiones distintas; sus transiciones detalladas requieren revisión. Los snapshots JSON solo conservan condiciones aplicadas, no reemplazan las FK principales ni autorizan contenido clínico.'),
 ('Selección exclusiva de Consulta','PRACTICA.base_inicial es una identidad reservada a las cuatro prácticas iniciales. La base CONSULTA es única por entidad e inmutable mediante edición común. Una personalizada tiene base nula aunque su nombre o clasificación sea Consulta. VERSION_OFERTA.portal_habilitado solo es válido para la base CONSULTA sin gestión obligatoria. Servidor valida la relación completa al reservar y reprogramar.'),
 ('Cobertura y atención particular','La ficha admite ninguna, una o varias afiliaciones; sin dato es distinto de sin cobertura. Cada evaluación y versión de turno conserva modalidad elegida y regla aplicada. El plan pertenece a su cobertura y la autorización a la afiliación/práctica correctas. No pasar a particular ante incompatibilidad sin elección expresa. Validación electrónica con cobertura está fuera del alcance.'),
 ('Confirmación y avisos','Un correo se programa dentro de T-72/T-48; WhatsApp a T-36 si sigue pendiente; T-24 sin confirmar cancela y libera con evento y aviso interno consistentes. Entrega, lectura, clic y texto libre nunca confirman. TOKEN_TURNO es limitado a versión y destinatario. Tareas antiguas se invalidan al cambiar horario; excepciones y turnos cercanos siguen D07.'),
 ('Reprogramación del profesional','LOTE_REPROGRAMACION e INTERVALO_LOTE identifican agenda y tramos. TURNO_AFECTADO deduplica el original y sigue aviso y reemplazo por separado. Todos los afectados vigentes reciben WhatsApp, incluidos confirmados, sin esperar T-36. No se libera la indisponibilidad ni se considera resuelto por enviar. Estado del original, liberación, cierre del lote y reemplazo siguen D07/D09.'),
 ('Continuidad del fichaje','TRANSICION_PROFESIONAL y CORRESPONDENCIA_FICHAJE son una alternativa de modelado, no una decisión sobre copia o proyección de fichas. El mecanismo autorizado debe preservar procedencia, pacientes compartidos y fichaje del profesional sin revelar operaciones ajenas. No se cambian tenant_id en bloque. Admisión, representación, convenios, chat, delegaciones y numeración del destino se resuelven expresamente en D04/D12.'),
 ('Eliminación y referencias históricas','PETICION_ELIMINACION registra la petición expresa verificada. No se propone ON DELETE CASCADE sobre fichas, turnos o auditoría. Retirar credenciales, cuenta y vínculos exige definir referencias históricas y anonimización antes del SQL; no se decide aquí retener una cuenta borrada como si siguiera activa. La baja institucional no activa eliminación ni cancela turnos.'),
 ('Operación y conservación','Las tablas de respaldo, restauración e incidente son metadatos de operación y pueden residir fuera de la base de negocio. No incluyen copias de datos ni secretos. Retención, índices físicos, RLS, normalización de correo, presupuesto y pruebas con rol real siguen pendientes. Los reportes básicos derivan de turnos/eventos y no requieren una tabla por reporte.'),
 ('Evolución fuera de beta','No se incorporan tablas de odontograma, prestaciones clínicas, facturación, cobros, suscripciones, pagadores, directorio público, IA, adjuntos, llamadas, calendarios externos o lista de espera automática. Suscripciones y tratamiento de datos por fin comercial se diseñan en I2/I3. SOLICITUD_ENTIDAD queda identificada como propuesta de docs/20 y no requisito aprobado.'),
]

TRACE=[
 ('R01 R01.a','ENTIDAD, SEDE, PACIENTE, PACIENTE_PROFESIONAL, CONTADOR_PACIENTE','CU-07,09,13-15,29-32,36-37; MP-07,35,37-41; D11 D12'),
 ('R02 R02.a R03','PRE_RESERVA, TURNO, OCUPACION_GLOBAL, IDEMPOTENCIA, POLITICA_OPERATIVA, CASO_ABUSO','CU-13-15,19-21,27,36-37; S03 S16; D01 D02 D15'),
 ('R04','MEMBRESIA, ROL, PERMISO, DELEGACION_AGENDA, DELEGACION_PERMISO','CU-04,08,11,15,17,19-22,28; MP-12..15,19..25,36; D03'),
 ('R05 R10','PROFESIONAL_ENTIDAD, TRANSICION_PROFESIONAL, CORRESPONDENCIA_FICHAJE, PETICION_ELIMINACION','CU-29-31; MP-37..41; ADR-011; D04 D12'),
 ('R06','PRACTICA, VERSION_PRACTICA, OFERTA, VERSION_OFERTA, HORARIO_AGENDA, RECURSO','CU-10-16,20,36-37; CAT-01..12; ADR-009; D05'),
 ('R07','COBERTURA_MEDICA, PLAN_COBERTURA, AFILIACION_PACIENTE, CONVENIO, EVALUACION_ELEGIBILIDAD','CU-09,12-15,20,36-37; D06; docs/18'),
 ('R08','VERSION_TURNO, CONFIRMACION_ASISTENCIA, EVENTO_TURNO, TOKEN_TURNO, TAREA_PERSISTIDA, AVISO_INTERNO','CU-18-22,24-26; OP-01/02; ADR-004; D07'),
 ('R09 R09.a','REPRESENTACION, REPRESENTACION_PERMISO, SOLICITUD_ADMISION, ACCESO_PACIENTE, ALCANCE_ACCESO','CU-01-04,06,09,14-19,23,28; S09 S15; D08 D14'),
 ('R10.a','CUENTA, CREDENCIAL_LOCAL, TOKEN_IDENTIDAD, SESION, FACTOR_MFA, INVITACION_EQUIPO','CU-05,07,35; ADR-012; S01 S06; D14 D16'),
 ('R11','CANAL_ENTIDAD, DESTINO_AVISO, AVISO_EXTERNO, INTENTO_ENVIO, EVENTO_PROVEEDOR, CONSUMO_CANAL','CU-24-26; docs/04; D09'),
 ('R11.a','CONVERSACION, PARTICIPANTE_CHAT, PARTICIPANTE_PERMISO, MENSAJE, LECTURA_MENSAJE','CU-16,23; MP-31/32; S14; D13'),
 ('R11.b','LOTE_REPROGRAMACION, INTERVALO_LOTE, TURNO_AFECTADO, SOLICITUD_GESTION','CU-21,26; OP-03; ADR-013; D07 D09'),
 ('R12 S10-S12','AUDITORIA_ENTIDAD, AUDITORIA_GLOBAL, ACCESO_SOPORTE, RESPALDO, ENSAYO_RESTAURACION, INCIDENTE','CU-32-34; OP-05; D11 D16'),
]

def field_text(f):
    prefix=(f['key']+' ') if f['key'] else ''
    return prefix+f['name']+(' ?' if f['null'] else '')+(' > '+f['ref'] if f['ref'] else ' : '+f['type'])

def wrap_line(txt,width,font='ER',size=13):
    words=txt.split(' '); lines=[]; s=''
    for word in words:
        t=(s+' '+word).strip()
        if s and pdfmetrics.stringWidth(t,font,size)>width:
            lines.append(s);s=word
        else:s=t
    if s:lines.append(s)
    return lines

pdfmetrics.registerFont(TTFont('ER',str(Path('C:/Windows/Fonts/arial.ttf'))))
pdfmetrics.registerFont(TTFont('ERB',str(Path('C:/Windows/Fonts/arialbd.ttf'))))

def pages():
    result=[]
    for group,names in GROUPS.items():
        # Máximo ocho entidades por lámina; todos sus campos se mantienen visibles.
        for offset in range(0,len(names),8):
            suffix=f' parte {offset//8+1}' if len(names)>8 else ''
            result.append((group+suffix,names[offset:offset+8]))
    return result

def make_diagram(title,names,index):
    cols=min(4,len(names)); rows=math.ceil(len(names)/cols)
    bw=340; gap=80; left=34; top=125
    heights=[]
    for n in names:
        heights.append(52+sum(max(1,len(wrap_line(field_text(f),bw-20))) for f in E[n]['fields'])*18+20)
    rowh=[max(heights[i*cols:(i+1)*cols]) for i in range(rows)]
    W=max(980,left*2+cols*bw+(cols-1)*gap)
    H=top+sum(rowh)+(rows-1)*105+95
    drawing=Drawing(W,H)
    drawing.add(Rect(0,0,W,H,fillColor=HexColor('#ffffff'),strokeColor=None))
    def toptext(value,y,size,font='ER'):
        size=min(size,(W-2*left)/pdfmetrics.stringWidth(value,font,1))
        drawing.add(String(left,H-y,value,fontName=font,fontSize=size,fillColor=HexColor('#000000' if font=='ERB' else '#43576a')))
    toptext(title,35,27,'ERB')
    toptext('PK UUIDv7  |  FK UUID  |  ? opcional  |  G global  T entidad  X excepción  |  ADR-015',64,16)
    toptext('Relaciones externas indicadas en cada FK; el registro completo resuelve cardinalidades y unicidades.',89,15)
    pos={}
    for k,n in enumerate(names):
        r,c=divmod(k,cols); y=H-top-sum(rowh[:r])-r*105-heights[k]
        pos[n]=(left+c*(bw+gap),y,bw,heights[k])
    # Conectores primero: líneas ortogonales entre tarjetas con extremos explícitos.
    local=[r for r in R if r['parent'] in names and r['child'] in names and r['field']!='tenant_id' and r['parent']!=r['child']]
    for k,r in enumerate(local):
        ax,ay,aw,ah=pos[r['parent']]; bx,by,bb,bh=pos[r['child']]
        if abs(ax-bx)>1:
            right=bx>ax; sx=ax+aw if right else ax; sy=ay+ah-24-(k%5)*13
            tx=bx if right else bx+bb; ty=by+bh-24-(k%5)*13
            mx=(sx+tx)/2+(k%3-1)*9
            pts=[(sx,sy),(mx,sy),(mx,ty),(tx,ty)]
        else:
            sx=ax+aw/2+(k%4)*13;sy=ay if by<ay else ay+ah
            tx=bx+bb/2+(k%4)*13;ty=by+bh if by<ay else by
            my=(sy+ty)/2;pts=[(sx,sy),(sx,my),(tx,my),(tx,ty)]
        for p,q in zip(pts,pts[1:]):drawing.add(Line(*p,*q,strokeColor=HexColor('#9aabba'),strokeWidth=1.25))
        # Los símbolos se dibujan en el exterior; FK identifica inequívocamente el rol.
        drawing.add(String(sx+3,sy+4,r['parent_card'],fontName='ER',fontSize=12,fillColor=HexColor('#354c61')))
        # Etiqueta del extremo hijo fuera del borde para que no quede tapada.
        label=r['child_card']
        lx=tx-7 if tx<=bx else tx+7
        drawing.add(String(lx,ty-5,label,fontName='ER',fontSize=12,textAnchor='end' if tx<=bx else 'start',fillColor=HexColor('#354c61')))
    for n,(x,y,w,h) in pos.items():
        e=E[n]; fill={'G':'#e6eef8','T':'#e4f2ee','X':'#fff0d5'}[e['scope']]
        drawing.add(Rect(x,y,w,h,fillColor=HexColor('#ffffff'),strokeColor=HexColor('#6f899a'),strokeWidth=1.4))
        drawing.add(Rect(x,y+h-43,w,43,fillColor=HexColor(fill),strokeColor=HexColor('#6f899a'),strokeWidth=1.0))
        size=min(17, (w-17)/max(pdfmetrics.stringWidth(n+' ['+e['scope']+']','ERB',1),1))
        drawing.add(String(x+9,y+h-27,n+' ['+e['scope']+']',fontName='ERB',fontSize=size,fillColor=HexColor('#000000')))
        yy=y+h-63
        for f in e['fields']:
            for line in wrap_line(field_text(f),w-20):
                drawing.add(String(x+10,yy,line,fontName='ER',fontSize=13,fillColor=HexColor('#183347'))); yy-=18
    toptext('* 0..N salvo UQ específica: consultar reglas de entidad y registro de relaciones. Modelo lógico propuesto.',H-32,14)
    stem=f'{index:02d}-'+re.sub('[^a-z0-9]+','-',title.lower()).strip('-')
    target=carpeta_lamina(OUT,stem)
    target.mkdir(parents=True,exist_ok=True)
    renderSVG.drawToFile(drawing,str(target/'diagrama.svg'))
    # PDF de lámina para rasterizar sin software de diagramación externo.
    pw,ph=landscape(A3); c=canvas.Canvas(str(TMP/(stem+'.pdf')),pagesize=(pw,ph))
    scale=min((pw-44)/W,(ph-44)/H)
    c.translate((pw-W*scale)/2,(ph-H*scale)/2);c.scale(scale,scale)
    renderPDF.draw(drawing,c,0,0);c.save()
    return dict(title=title,names=names,positions=pos,w=W,h=H,stem=stem,drawing=drawing)

def mxvertex(root,ident,value,x,y,w,h,style,parent='1'):
    c=ET.SubElement(root,'mxCell',id=ident,value=value,style=style,vertex='1',parent=parent)
    ET.SubElement(c,'mxGeometry',x=str(x),y=str(y),width=str(w),height=str(h),attrib={'as':'geometry'})
    return c

def drawio_page(mxfile,d):
    dia=ET.SubElement(mxfile,'diagram',id=d['stem'],name=d['title'])
    graph=ET.SubElement(dia,'mxGraphModel',dx='1600',dy='1100',grid='1',page='1',pageWidth=str(d['w']),pageHeight=str(d['h']))
    root=ET.SubElement(graph,'root');ET.SubElement(root,'mxCell',id='0');ET.SubElement(root,'mxCell',id='1',parent='0')
    mxvertex(root,'title',d['title']+'\nPK UUIDv7 | FK UUID | ? opcional | G global T entidad X excepción',25,15,d['w']-50,82,'text;html=0;align=left;fontSize=24;whiteSpace=wrap;')
    for n,(x,y,w,h) in d['positions'].items():
        e=E[n]; y=d['h']-y-h; fill={'G':'#e6eef8','T':'#e4f2ee','X':'#fff0d5'}[e['scope']]
        mxvertex(root,n,n+' ['+e['scope']+']',x,y,w,h,'swimlane;html=0;startSize=43;horizontal=1;rounded=0;fillColor='+fill+';swimlaneFillColor=#ffffff;strokeColor=#6f899a;fontSize=17;fontStyle=1;')
        content='\n'.join('\n'.join(wrap_line(field_text(f),w-20)) for f in e['fields'])
        mxvertex(root,n+'_fields',content,7,48,w-14,h-51,'text;html=0;align=left;verticalAlign=top;fontSize=13;fontStyle=0;whiteSpace=wrap;spacing=3;',n)
    for k,r in enumerate(R):
        if r['child'] not in d['names'] or r['parent'] not in d['names'] or (r['field']=='tenant_id' and not d.get('integral')):continue
        cell=ET.SubElement(root,'mxCell',id='r'+str(k),value=r['parent_card']+' / '+r['child_card']+' · '+r['field'],edge='1',parent='1',source=r['parent'],target=r['child'],style='edgeStyle=orthogonalEdgeStyle;rounded=0;html=0;endArrow=none;startArrow=none;strokeColor=#7f96a8;fontSize=12;labelBackgroundColor=#ffffff;')
        ET.SubElement(cell,'mxGeometry',relative='1',attrib={'as':'geometry'})
    mxvertex(root,'legend','Relaciones externas: destino en la FK. * UQ puede reducir hijos a 0..1. Consultar diccionario y registro completo.',25,d['h']-60,d['w']-50,38,'text;html=0;align=left;fontSize=15;')

def mermaid(names):
    lines=['erDiagram']
    for n in names:
        lines.append('    '+n+' {')
        for f in E[n]['fields']:
            key=' '+f['key'] if f['key'] else ''
            comment=('opcional ' if f['null'] else 'obligatorio ')+(('ref '+f['ref']) if f['ref'] else f['note'])
            lines.append('        '+f['type']+' '+f['name']+key+' "'+comment.replace('"',"'")+'"')
        lines.append('    }')
    for r in R:
        if r['child'] in names and r['parent'] in names:
            p='|o' if r['parent_card']=='0..1' else '||'
            child='o|' if r['child_card']=='0..1' else 'o{'
            lines.append('    '+r['parent']+' '+p+'--'+child+' '+r['child']+' : "'+r['field']+'"')
    return '\n'.join(lines)+'\n'

def write_sources(ds):
    model=dict(title='Modelo entidad relación de gestión administrativa de salud',version='1.1',date='2026-10-17',status='Modelo propuesto para revisión; identificadores aceptados en ADR-015',identifier_policy=dict(decision='ADR-015',primary_key_type='uuid',primary_key_version=7,primary_key_default='uuidv7()',foreign_key_type='uuid',foreign_keys_generated=False,visible_patient_number_type='entero',visible_patient_number_scope=['tenant_id','profesional_id']),entities=list(E.values()),relationships=R,notes=NOTES,traceability=TRACE)
    (OUT/'modelo-er.json').write_text(json.dumps(model,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'modelo-completo.mmd').write_text(mermaid(list(E)),encoding='utf-8')
    mxfile=ET.Element('mxfile',host='app.diagrams.net',version='24.7.17',compressed='false')
    # Plano integral con todas las entidades y conectores entre áreas.
    # Es un lienzo grande para zoom; las láminas siguientes sirven para lectura.
    allnames=list(E); positions={}; yy=125
    for start in range(0,len(allnames),4):
        row=allnames[start:start+4]; heights=[]
        for n in row:
            heights.append(72+18*sum(len(wrap_line(field_text(f),320)) for f in E[n]['fields']))
        for j,n in enumerate(row):positions[n]=(34+j*460,yy,340,heights[j])
        yy+=max(heights)+120
    master=dict(stem='00-modelo-integral',title='00 Modelo integral con relaciones entre áreas',names=allnames,positions={},w=1840,h=yy+100,integral=True)
    master['positions']={n:(x,master['h']-y-h,w,h) for n,(x,y,w,h) in positions.items()}
    drawio_page(mxfile,master)
    for d in ds:
        drawio_page(mxfile,d)
        one=ET.Element('mxfile',host='app.diagrams.net',compressed='false');drawio_page(one,d)
        target=carpeta_lamina(OUT,d['stem'])
        target.mkdir(parents=True,exist_ok=True)
        ET.ElementTree(one).write(target/'diagrama.drawio',encoding='utf-8',xml_declaration=True)
        (target/'diagrama.mmd').write_text(mermaid(d['names']),encoding='utf-8')
    ET.ElementTree(mxfile).write(OUT/'modelo-er-completo.drawio',encoding='utf-8',xml_declaration=True)
    md=['# Modelo entidad relación completo','', '17 de octubre de 2026 · Versión 1.1 · Modelo lógico propuesto para revisión con identificadores aceptados en ADR-015','',f'{len(E)} entidades, {len(R)} referencias FK y {len(ds)} láminas de detalle, más un plano integral en draw.io con conectores entre áreas.','']
    for title,body in NOTES:md += ['## '+title,'',body,'']
    md+=['## Documentación explicativa paso a paso','', 'La [colección de estudio](arquitectura/modelo-er/partes/README.md) organiza las 17 láminas en carpetas numeradas, con PDF explicativos de todas las entidades, campos, reglas y relaciones. Empezar por el [índice y ruta de lectura](arquitectura/modelo-er/partes/00-indice-y-ruta-de-lectura.pdf). Regenerar con `scripts/generar-guia-modelo-er.py` tras actualizar el modelo. Los ejemplos son ficticios y no aprueban decisiones pendientes.','']
    md+=['## Láminas editables','']
    for d in ds:
        relative=carpeta_lamina(OUT,d['stem']).relative_to(OUT).as_posix()
        md += [f'- [{d["title"]}](arquitectura/modelo-er/{relative}/README.md) · [draw.io](arquitectura/modelo-er/{relative}/diagrama.drawio) · [SVG](arquitectura/modelo-er/{relative}/diagrama.svg)']
    md+=['','## Diccionario de datos','']
    for n,e in E.items():
        md+=['### '+n,'',e['purpose'],'',f'Ámbito {e["scope"]}. Estado: {e["status"]}. Fuentes: {e["source"]}.','', '| Campo | Tipo | Clave o destino | Nulo |', '| --- | --- | --- | --- |']
        for f in e['fields']:md += [f'| {f["name"]} | {TYPE_LABEL[f["type"]]} | {f["key"]} {f["ref"] or ""} | {"Sí" if f["null"] else "No"} |']
        md+=['',e['rules'] or 'Claves y ámbito sujetos a las convenciones generales.','']
    md+=['## Registro completo de relaciones','','Cada fila expresa la cantidad de padres por hijo y la colección inversa. Las unicidades específicas reducen el máximo de hijos; no son concesiones de acceso.','', '| Hijo y FK | Padre | Padre por hijo | Hijos por padre | Integridad |','| --- | --- | --- | --- | --- |']
    for r in R:md += [f'| {r["child"]}.{r["field"]} | {r["parent"]} | {r["parent_card"]} | {r["child_card"]} | {r["integrity"]} |']
    md+=['','## Trazabilidad funcional','', '| Regla | Entidades principales | Casos y decisiones |','| --- | --- | --- |']
    for t in TRACE:md+=['| '+' | '.join(t)+' |']
    md+=['','## Fuentes documentales','', 'Prevalecen las reglas y ADR vigentes. Los análisis históricos sustituidos no se usan como norma. El inventario conserva todos los documentos del relevamiento y permite rastrear el corte utilizado.','']
    for p in sorted((ROOT/'docs').glob('*.md')):
        if p.name.startswith('23-'):continue
        md += [f'- [{p.name}]({p.name})']
    for p in sorted((ROOT/'docs/arquitectura/adr').glob('*.md')):md += [f'- [{p.name}](arquitectura/adr/{p.name})']
    (ROOT/'docs/23-modelo-entidad-relacion.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    (OUT/'README.md').write_text('''# Modelo entidad relación editable

Abrir `modelo-er-completo.drawio` en diagrams.net o draw.io. Contiene todas las láminas con entidades, atributos y conectores nativos; no son capturas incrustadas. También se entregan láminas individuales.

## Explicación paso a paso

La [colección de estudio](partes/README.md) organiza las 17 láminas en carpetas numeradas. Cada carpeta incluye una explicación detallada en PDF y Markdown, todos los campos y relaciones, ejemplos ficticios y su diagrama en PDF, SVG, draw.io y Mermaid. Comenzar por el [índice y ruta de lectura](partes/00-indice-y-ruta-de-lectura.pdf) y la [explicación de identidad y credenciales, parte 1](partes/01-identidad-y-credenciales-parte-1/explicacion.pdf).

Contenido pedagógico en el `contenido.json` de cada parte y en `partes/glosario-campos.json`; regenerar la colección con `scripts/generar-guia-modelo-er.py` después de actualizar las fuentes del modelo. El material explicativo conserva el estado propuesto y los pendientes de las decisiones originales.

`modelo-er.json` es la fuente estructurada generada por `scripts/generar-modelo-er.py`. El generador mantiene fuentes, Word, diagramas SVG, Mermaid y PDF de láminas coherentes. Los PNG de trabajo quedan en `tmp/modelo-er` y no son entregables.

El documento principal es `docs/23-modelo-entidad-relacion.md`; Word en `output/docx/modelo-entidad-relacion-salud.docx` y PDF en `output/pdf/modelo-entidad-relacion-salud.pdf`. Versión 1.1, 17 de octubre de 2026. [ADR-015](../adr/015-identificadores-uuidv7.md) fija UUIDv7 para PK internas y tipo uuid para sus FK, incluido tenant_id; el número visible del paciente sigue siendo entero por profesional y entidad. No es SQL ejecutable.

G global restringido, T entidad con tenant_id, X excepción administrativa. Las líneas de cada lámina conectan su propio grupo. Todas las referencias externas están escritas en la FK y enumeradas en el registro completo del documento. 1/0..1 expresa padres por hijo; 0..N salvo UQ expresa hijos por padre. Los puentes resuelven N a M. La tabla ENTIDAD es raíz global de aislamiento y sus referencias tenant_id se omiten como líneas para evitar saturación.

El modelo respeta las decisiones confirmadas y marca propuestas y pendientes. No cierra D01-D16, ni implementa permisos, RLS, capacidad, canales o eliminación. Las entidades X documentan una alternativa excepcional aún por revisar, sin traslado automático de fichas.

Las ediciones manuales del draw.io se preservan al ejecutar el generador salvo `--sobrescribir-editables`. Para conservarlas como fuente permanente, trasladar su intención al generador y regenerar bajo revisión.
''',encoding='utf-8')
    # Inventario para auditoría del relevamiento, sin copiar secretos ni documentos externos.
    import hashlib
    sources=[]
    for p in sorted((ROOT/'docs').rglob('*.md')):
        if p.name.startswith('23-') or 'modelo-er' in p.parts:continue
        sources.append(dict(path=p.relative_to(ROOT).as_posix(),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),vigente='archivo' not in p.parts))
    (OUT/'fuentes-documentales.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding='utf-8')

def style_table(t,widths=None):
    t.autofit=False
    for i,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            if widths:cell.width=Mm(widths[j])
            tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
            tcPr.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCEAF2' if i==0 else ('F4F7F9' if i%2==0 else 'FFFFFF'));tcPr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                v=OxmlElement('w:'+edge);v.set(qn('w:w'),'85');v.set(qn('w:type'),'dxa');margins.append(v)
            tcPr.append(margins)
            cell.vertical_alignment=1
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                # El encabezado debe permanecer junto a la primera fila de datos.
                if i==0:p.paragraph_format.keep_with_next=True
                for run in p.runs:run.font.size=Pt(8);run.bold=i==0;run.font.color.rgb=RGBColor(0,0,0)
        trpr=row._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trpr.append(el)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)

def build_docx(ds):
    doc=Document();s=doc.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(18);s.left_margin=s.right_margin=Mm(18)
    for sn in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[sn];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    # Eliminar reglas de borde heredadas de la plantilla de Word.
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
    doc.styles['Heading 2'].paragraph_format.space_after=Pt(4)
    for sn,size in [('Title',26),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:doc.styles[sn].font.size=Pt(size)
    doc.add_paragraph('Modelo entidad relación',style='Title')
    doc.add_paragraph('Sistema de Gestión Administrativa de Salud',style='Subtitle')
    doc.add_paragraph('Modelo lógico completo y diccionario de datos\n17 de octubre de 2026 · Versión 1.1')
    doc.add_paragraph(f'El modelo organiza {len(E)} entidades y {len(R)} referencias para revisar el diseño de pacientes, permisos, agendas, turnos y comunicaciones antes de construir la beta. Se entrega en Word editable y en draw.io con objetos nativos.')
    doc.add_paragraph('Las reglas confirmadas conservan su vigencia. Las tablas y atributos que las desarrollan son una propuesta de diseño; la continuidad del fichaje, los permisos pendientes y los parámetros abiertos requieren revisión. No hay base de datos implementada ni migración SQL en esta entrega.')
    doc.add_heading('Guía de consulta',1)
    doc.add_paragraph('Primero se explican alcance, claves y restricciones. Luego se presentan las láminas por área y el diccionario completo. El registro de relaciones permite seguir toda FK, incluso entre láminas. Al final figuran trazabilidad y fuentes.')
    doc.add_heading('Contenido por área',2)
    for g,names in GROUPS.items():doc.add_paragraph(g+' · '+str(len(names))+' entidades')
    doc.add_page_break()
    doc.add_heading('Criterios del modelo',1)
    for title,body in NOTES:
        doc.add_heading(title,2);doc.add_paragraph(body)
    # Láminas con los mismos atributos nativos de draw.io; imágenes para lectura.
    for d in ds:
        sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Mm(420);sec.page_height=Mm(297);sec.left_margin=sec.right_margin=Mm(12);sec.top_margin=sec.bottom_margin=Mm(10)
        img=TMP/(d['stem']+'.png')
        # Ajuste simultáneo ancho y alto al área imprimible A3.
        im=Image.open(img);w,h=im.size; maxw=396;maxh=264; scale=min(maxw/w,maxh/h)
        doc.add_picture(str(img),width=Mm(w*scale),height=Mm(h*scale))
        doc.paragraphs[-1].paragraph_format.space_after=Pt(0)
        doc.paragraphs[-1].paragraph_format.line_spacing=1
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.PORTRAIT;sec.page_width=Mm(210);sec.page_height=Mm(297);sec.left_margin=sec.right_margin=Mm(18);sec.top_margin=sec.bottom_margin=Mm(18)
    doc.add_heading('Diccionario de datos',1)
    doc.add_paragraph('UUID identifica las claves internas: nuevas PK UUIDv7 generadas por PostgreSQL y FK del mismo tipo uuid que conservan el valor del padre, según ADR-015. El número visible del paciente sigue siendo entero por profesional y entidad. Instante conserva el punto temporal y su zona de presentación. Enum requiere catálogo de valores y transiciones aprobado. Estructura versionada conserva condiciones administrativas; no incorpora información clínica. Nulabilidad y campos personales son propuestas sujetas a revisión.')
    doc.add_paragraph('En toda tabla T se añade tenant_id obligatorio, FK a ENTIDAD y parte de cada referencia local compuesta. La clave primaria es id y existe UQ(tenant_id,id). Se omite repetir tenant_id en las tablas del diccionario para facilitar la lectura; aparece en todas las láminas y fuentes estructuradas.')
    for group,names in GROUPS.items():
        doc.add_heading(group,1)
        for n in names:
            e=E[n];doc.add_heading(n,2)
            p=doc.add_paragraph(e['purpose']);p.paragraph_format.keep_with_next=True
            p=doc.add_paragraph(f'Ámbito {e["scope"]} · {e["status"]}\nFuentes {e["source"]}');p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
            fields=[f for f in e['fields'] if f['name']!='tenant_id']
            t=doc.add_table(rows=1,cols=4)
            for c,v in zip(t.rows[0].cells,['Campo','Tipo','Clave y destino','Nulo']):c.text=v
            for f in fields:
                row=t.add_row().cells
                vals=[f['name'],TYPE_LABEL[f['type']],f['key']+(' > '+f['ref'] if f['ref'] else ''),'Sí' if f['null'] else 'No']
                for c,v in zip(row,vals):c.text=v
            style_table(t,[57,36,66,15])
            p=doc.add_paragraph(e['rules'] or 'Se aplican las convenciones generales de claves, ámbito y autorización.');p.paragraph_format.space_before=Pt(5)
    doc.add_heading('Registro completo de relaciones',1)
    doc.add_paragraph('Cada FK refiere a un padre. Una FK obligatoria requiere exactamente uno; una opcional permite cero o uno. El padre admite cero o muchos hijos, salvo la unicidad indicada en las reglas de la entidad. Las relaciones N a M se descomponen mediante tablas puente. Este registro incluye todos los extremos externos y la raíz de aislamiento.')
    for group,names in GROUPS.items():
        doc.add_heading(group,2)
        t=doc.add_table(rows=1,cols=4)
        for c,v in zip(t.rows[0].cells,['Hijo y campo FK','Padre','Padres por hijo','Hijos por padre']):c.text=v
        for r in R:
            if r['child'] not in names:continue
            row=t.add_row().cells
            for c,v in zip(row,[r['child']+'\n'+r['field'],r['parent'],r['parent_card'],r['child_card']]):c.text=v
        style_table(t,[72,60,21,21])
    doc.add_heading('Trazabilidad y decisiones abiertas',1)
    for rule,entities,refs in TRACE:
        doc.add_heading(rule.replace('.',' '),2);doc.add_paragraph(entities);doc.add_paragraph(refs)
    doc.add_paragraph('D01 a D16 conservan su estado documental. No se fijan TTL, cuotas, tiempos de prácticas, transiciones detalladas, proveedores ni conservación por generar el modelo. D10 y el directorio público permanecen fuera de beta. La futura revisión debe validar SQL, índices, claves compuestas, RLS, concurrencia y compatibilidad de estados con el rol real de aplicación.')
    doc.add_heading('Fuentes del relevamiento',1)
    doc.add_paragraph('Las reglas R01 a R12 y ADR vigentes son la autoridad. El relevamiento inicial revisó los documentos 01 a 22, el glosario, estado consolidado y matriz de permisos. ADR-015 registra la política de identificadores aceptada el 17 de octubre y las fuentes técnicas consultadas para esa decisión. Los análisis sustituidos y las tarifas históricas no agregan requisitos al modelo.')
    for p in sorted((ROOT/'docs').glob('*.md')):
        if p.name.startswith('23-'):continue
        doc.add_paragraph('docs/'+p.name)
    for p in sorted((ROOT/'docs/arquitectura/adr').glob('*.md')):doc.add_paragraph('docs/arquitectura/adr/'+p.name)
    doc.core_properties.title='Modelo entidad relación del Sistema de Gestión Administrativa de Salud'
    doc.core_properties.subject='Modelo lógico propuesto y diccionario de datos'
    doc.core_properties.author='Proyecto Sistema de Gestión Administrativa de Salud'
    # Las secciones enlazan al mismo pie; escribirlo una sola vez.
    footer=doc.sections[0].footer.paragraphs[0];footer.clear();footer.alignment=2
    footer.add_run('Modelo ER · 17 octubre 2026 · ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for run in footer.runs:run.font.size=Pt(8)
    path=ROOT/'output/docx/modelo-entidad-relacion-salud.docx';doc.save(path)
    return path

def main():
    import argparse, subprocess, shutil
    p=argparse.ArgumentParser();p.add_argument('--sobrescribir-editables',action='store_true');args=p.parse_args()
    # Conservar ediciones draw.io existentes salvo autorización explícita por flag.
    saved={f.relative_to(OUT):f.read_bytes() for f in OUT.rglob('*.drawio')} if not args.sobrescribir_editables else {}
    ds=[make_diagram(title,names,i+1) for i,(title,names) in enumerate(pages())]
    write_sources(ds)
    for n,data in saved.items():
        # Actualización semántica acotada: conservar geometría y ediciones manuales.
        # Las FK muestran su destino; su tipo uuid se define en la política común.
        tree=ET.fromstring(data)
        for cell in tree.iter('mxCell'):
            if cell.get('id')=='title':
                cell.set('value',cell.get('value','').replace('PK clave | FK referencia','PK UUIDv7 | FK UUID'))
            if cell.get('id','').endswith('_fields'):
                value=cell.get('value','')
                cell.set('value',re.sub(r'(?<= : )id(?=<|\n|$)','uuid',value))
        (OUT/n).write_bytes(ET.tostring(tree,encoding='utf-8',xml_declaration=True))
    poppler=Path('C:/Users/gerag/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe')
    for d in ds:
        subprocess.run([str(poppler),'-scale-to','2300','-png','-singlefile',str(TMP/(d['stem']+'.pdf')),str(TMP/d['stem'])],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
    path=build_docx(ds)
    print(json.dumps(dict(entities=len(E),relationships=len(R),diagrams=len(ds),docx=str(path)),ensure_ascii=False))

if __name__=='__main__':main()
